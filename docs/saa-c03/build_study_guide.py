"""
Build the SAA-C03 practical-lab Word document.

Experiment chapters are added only after the learner provides notes or
screenshots. This first build creates the professional framework, exam map,
intake form for Experiment 1, glossary, and appendices.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from docx_theme import (  # noqa: E402
    BODY,
    NAVY,
    SLATE,
    TEAL,
    add_body,
    add_bullet,
    add_callout,
    add_figure,
    add_header_and_footer,
    add_mixed_paragraph,
    add_numbered,
    add_page_break,
    add_styled_table,
    add_toc,
    patch_docx_update_fields,
    set_document_defaults,
    set_picture_alt_text,
    set_run_font,
    set_update_fields_on_open,
)
from generate_assets import generate_all  # noqa: E402
from registry import DOCUMENT, EXPERIMENTS  # noqa: E402


OUTPUT = ROOT / "AWS-SAA-C03-Practical-Experiments-and-Study-Guide.docx"


def heading(document: Document, text: str, level: int = 1):
    return document.add_heading(text, level=level)


def build_cover(document: Document, assets: dict[str, Path]) -> None:
    banner = document.add_paragraph()
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    banner.paragraph_format.space_after = Pt(10)
    banner.paragraph_format.space_before = Pt(0)
    run = banner.add_run()
    cover_shape = run.add_picture(str(assets["cover"]), width=Inches(6.5))
    set_picture_alt_text(
        cover_shape,
        "Navy cover banner for the SAA-C03 practical-lab journal, showing the document title, learner name Kritheebhan, version 0.1, and status awaiting Experiment 1 notes and screenshots.",
        title="Document cover banner",
    )

    for line, size, color, bold, space in [
        ("A living lab journal for hands-on AWS practice", 14, TEAL, False, 8),
        ("Prepared for " + DOCUMENT["learner"], 16, NAVY, True, 4),
        ("Software engineer  ·  SAA-C03 exam preparation", 12, SLATE, False, 14),
    ]:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space)
        run = p.add_run(line)
        set_run_font(run, size, bold=bold, color=color)

    add_styled_table(
        document,
        ["Item", "Value"],
        [
            ["Document title", DOCUMENT["title"]],
            ["Exam code", "SAA-C03"],
            ["Document version", DOCUMENT["version"] + " — " + DOCUMENT["edition"]],
            ["Date", DOCUMENT["date"]],
            ["Current status", DOCUMENT["status"]],
            ["Experiments completed", "0"],
            ["Next checkpoint", "After Experiment 5"],
            ["Classification", "Personal study notes. Not official AWS training."],
        ],
        col_widths=[2.1, 4.4],
    )

    add_callout(
        document,
        "next",
        "This first edition is the document framework only. Experiment 1 will be written after you send the experiment name, course section, notes, AWS Region, screenshots, and any errors you saw. No lab steps will be invented.",
        title="Start here",
    )


def build_inside_cover(document: Document) -> None:
    heading(document, "About this document", 1)
    add_body(
        document,
        "This study guide records practical AWS experiments in a consistent, exam-aware format. "
        "It is written for a beginner-friendly reader who wants to understand why each console action matters, not only which button to click.",
    )
    add_body(
        document,
        "The document is a personal lab journal. It is not an AWS publication, not a substitute for the official exam guide, and not a copy of any paid course. "
        "Course videos, slide decks, and vendor lab scripts are rewritten in original wording. Where a technical value is missing or unclear, the text is marked [Verification required] instead of guessed.",
    )

    heading(document, "Purpose", 2)
    add_bullet(document, "Capture each hands-on experiment in a form that is useful later during exam revision.")
    add_bullet(document, "Map every lab to SAA-C03 domains, AWS services, and common exam traps.")
    add_bullet(document, "Keep screenshots, resource names, cleanup steps, and cost warnings together in one place.")
    add_bullet(document, "Protect account identifiers and other sensitive values before the notes are stored.")

    heading(document, "What this document is not", 2)
    add_bullet(document, "It is not official AWS training or an AWS exam dump.")
    add_bullet(document, "It is not a verbatim transcript of a course.")
    add_bullet(document, "It does not include live account numbers, access keys, passwords, email addresses, or public IP addresses.")
    add_bullet(document, "It does not invent missing lab details to make a chapter look complete.")

    heading(document, "Document control", 2)
    add_styled_table(
        document,
        ["Version", "Date", "Author", "Change summary"],
        [
            [
                "0.1",
                "26 August 2026",
                DOCUMENT["learner"] + " with study-guide editor",
                "Created the professional framework, SAA-C03 domain map, figure conventions, Experiment 1 intake form, glossary, and appendices. No experiment chapter written yet.",
            ]
        ],
        col_widths=[0.9, 1.4, 1.6, 2.6],
    )

    heading(document, "How this Word file is rebuilt", 2)
    add_body(
        document,
        "The .docx file is generated so headings, tables, callouts, and figure style stay consistent when later experiments are added. From the repository root, install the Python packages in docs/saa-c03/requirements.txt and run python3 docs/saa-c03/build_study_guide.py. "
        "Place redacted screenshots in docs/saa-c03/screenshots and lab notes in docs/saa-c03/notes. Do not edit experiment wording only inside Word if you also keep the generator; the next rebuild would replace the file.",
    )

    heading(document, "Sources used for the framework", 2)
    add_body(
        document,
        "Exam structure in this edition is based on the official AWS Certified Solutions Architect – Associate (SAA-C03) exam guide published by AWS, rewritten in original language. "
        "Hands-on steps will come only from your notes and screenshots. Pricing statements will be limited to official qualitative guidance unless a figure is later verified.",
    )
    add_callout(
        document,
        "verify",
        "AWS can update exam codes, domain weights, in-scope services, and pricing. Before you book the exam, confirm on the AWS Certification website that SAA-C03 is still the exam you intend to take and that the domain weights have not changed.",
    )


def build_how_to_use(document: Document, assets: dict[str, Path]) -> None:
    heading(document, "How to use this study guide", 1)
    add_body(
        document,
        "Use this file as a lab notebook first and as an exam companion second. Complete the hands-on work, then read the exam mapping, security notes, cost notes, and knowledge-check questions for that experiment.",
    )

    heading(document, "Suggested study rhythm", 2)
    add_numbered(document, "Perform the experiment in your own AWS account using least-privilege access.")
    add_numbered(document, "Capture screenshots as you go. Hide account IDs, emails, keys, and IP addresses before you share them.")
    add_numbered(document, "Send the notes package described in the Experiment 1 intake section.")
    add_numbered(document, "Review the drafted chapter. Confirm or correct any [Verification required] items.")
    add_numbered(document, "After the chapter is approved, complete the five knowledge-check questions without looking at the answers.")
    add_numbered(document, "Delete lab resources using the cleanup order in that chapter, then confirm that nothing billable remains.")

    heading(document, "Working process for every new experiment", 2)
    add_figure(
        document,
        assets["process"],
        "Figure 1: Process used to add each experiment to this document.",
        "Five-step process: learner sends notes and screenshots, missing items are listed, a chapter is drafted in original wording, technical values are confirmed, and the approved chapter is added to the index.",
    )
    add_numbered(document, "You provide the experiment name, course section, notes, AWS Region, screenshots, and errors encountered.")
    add_numbered(document, "The supplied material is reviewed. Missing screenshots or unclear values are listed instead of guessed.")
    add_numbered(document, "The experiment chapter is drafted using the standard 16-section template.")
    add_numbered(document, "You verify technical values, resource names, and any [Verification required] markers.")
    add_numbered(document, "The approved chapter is added to the main document.")
    add_numbered(document, "Terminology, headings, tables, callouts, and figure style stay consistent with earlier chapters.")
    add_numbered(document, "The table of contents, figure numbering, experiment index, service index, and glossary are updated.")
    add_numbered(document, "After every five experiments, a revision checkpoint and overall progress summary are added.")
    add_numbered(document, "Before a final edition, the full document is checked for technical consistency, image placement, spelling, grammar, layout, and cleanup instructions.")

    add_callout(
        document,
        "note",
        "An experiment chapter is not started until notes or screenshots for that experiment are available. The pages that follow are the framework only.",
    )


def build_conventions(document: Document, assets: dict[str, Path]) -> None:
    heading(document, "Document conventions", 1)
    add_body(
        document,
        "These rules keep later chapters readable and consistent. Future experiment text will follow this section even if course videos use different wording.",
    )

    heading(document, "Headings", 2)
    add_styled_table(
        document,
        ["Level", "Used for", "Appears in table of contents"],
        [
            ["Heading 1", "Major parts, such as exam map, an experiment title block, glossary, and appendices", "Yes"],
            ["Heading 2", "The 16 standard experiment sections and other major subsections", "Yes"],
            ["Heading 3", "Individual procedure steps and smaller subsections", "Yes"],
            ["Heading 4", "Rare sub-steps inside a long procedure", "No"],
        ],
        col_widths=[1.4, 3.6, 1.5],
    )

    heading(document, "Standard experiment structure", 2)
    add_body(document, "Every completed experiment uses the same 16 sections in this order:")
    sections = [
        "1. Experiment Overview",
        "2. SAA-C03 Exam Mapping",
        "3. Learning Objectives",
        "4. Architecture",
        "5. Prerequisites",
        "6. Resources and Configuration",
        "7. Step-by-Step Procedure",
        "8. Verification",
        "9. Troubleshooting",
        "10. Security Best Practices",
        "11. Cost Considerations",
        "12. Cleanup Procedure",
        "13. Important SAA-C03 Exam Points",
        "14. Key Terms",
        "15. Knowledge Check",
        "16. Experiment Summary",
    ]
    for item in sections:
        add_bullet(document, item)

    heading(document, "Callouts", 2)
    add_body(document, "Shaded boxes highlight information that should not be lost in a long procedure.")
    add_styled_table(
        document,
        ["Callout", "Meaning"],
        [
            ["Exam tip", "A fact, comparison, or trap that often appears in SAA-C03-style questions."],
            ["Security", "An access-control, encryption, or network-isolation practice."],
            ["Cost", "A chargeable resource, Free Tier caution, or cleanup reminder. No unverified prices."],
            ["Verification required", "A value that was missing, unreadable, or not confirmed from your notes."],
            ["Privacy", "A reminder to hide account data before sharing screenshots."],
            ["Next action", "What you need to send or confirm before the next document update."],
            ["Warning", "A step that can cause outage, data loss, or unexpected charges if done incorrectly."],
        ],
        col_widths=[2.0, 4.5],
    )

    heading(document, "Tables", 2)
    add_body(
        document,
        "Resource lists, troubleshooting, and key terms always use tables with navy headers and alternating row shading. "
        "Column names stay the same across experiments so you can scan quickly during revision.",
    )

    heading(document, "Figures and screenshots", 2)
    add_body(
        document,
        "Each image is placed immediately after the related instruction. A figure number and caption sit under the image. "
        "File names follow EXP##-Step##-short-description.png. Alternative text is stored on the image for accessibility.",
    )
    add_styled_table(
        document,
        ["Rule", "How it is applied"],
        [
            ["Placement", "Insert the screenshot in the matching procedure step, not in an appendix dump."],
            ["Size", "Fit the page width without shrinking text in the screenshot below a readable size."],
            ["Quality", "Keep the original image quality. Do not add compression artifacts."],
            ["Cropping", "Browser chrome may be cropped when it adds no technical value."],
            ["Markup", "Arrows or boxes are added only when they identify an important setting."],
            ["Values", "Important technical values visible in a screenshot are never edited."],
            ["Unclear images", "The caption will say [A clearer screenshot is required]."],
            ["Sensitive data", "Account IDs, emails, keys, and IP addresses are hidden or blurred."],
        ],
        col_widths=[1.7, 4.8],
    )

    heading(document, "Architecture diagram style", 2)
    add_figure(
        document,
        assets["legend"],
        "Figure 2: Architecture diagram conventions used in later experiment chapters.",
        "Legend diagram showing users, internet edge, a VPC boundary, a public subnet with load balancer and NAT, a private subnet with compute and data stores, security groups, IAM roles, encryption, and logging. This is a style key, not a deployed lab.",
    )
    add_body(
        document,
        "Each experiment architecture will show users or clients, the AWS services in use, the request or data flow, network boundaries, and security controls. A short explanation will follow the diagram. No experiment architecture is drawn until that lab’s notes exist.",
    )

    heading(document, "Terminology", 2)
    add_body(
        document,
        "The document uses consistent AWS terms. A virtual private cloud is written as VPC. "
        "An Availability Zone is written in full on first use in a chapter, then AZ. "
        "Identity and Access Management is written as IAM. "
        "Service names keep the Amazon or AWS prefix used in the official service list when the distinction matters, and a short name afterward.",
    )
    add_exam_tip = add_callout
    add_exam_tip(
        document,
        "exam",
        "The exam expects you to choose the service that fits the requirement, not the service that appeared in a single lab. Each chapter will call out similar services that are easy to confuse, such as security groups compared with network ACLs, or NAT gateways compared with internet gateways.",
    )


def build_privacy(document: Document) -> None:
    heading(document, "Privacy and redaction rules", 1)
    add_body(
        document,
        "Lab screenshots often contain identifiers that should not be stored in a shared study file. Apply these rules before an image is inserted.",
    )
    add_styled_table(
        document,
        ["Hide or blur this", "Examples", "Why"],
        [
            ["AWS account IDs", "12-digit account number in the console header or ARN", "Identifies the account and can appear in ARNs and billing views."],
            ["Email addresses", "IAM user sign-in, root email, SNS notifications", "Personal data and account recovery path."],
            ["Access keys and secrets", "AKIA… keys, secret key pairs, session tokens", "These grant programmatic access and must never be stored."],
            ["Passwords and MFA secrets", "Login forms, authenticator seeds", "Direct account takeover risk."],
            ["IP addresses", "Elastic IPs, public instance IPs, on-premises CIDRs that identify you", "Can expose running systems."],
            ["Customer or employer names", "Bucket names or tags that identify a workplace", "Keeps this journal personal and reusable."],
            ["Private keys and certificates", ".pem contents, certificate bodies", "Allow SSH or TLS impersonation."],
        ],
        col_widths=[1.8, 2.4, 2.3],
    )
    add_callout(
        document,
        "privacy",
        "If a screenshot still shows a 12-digit account ID, an email address, an access key, or a public IP, do not treat the chapter as final. Send a redacted replacement image. Resource names such as saa-lab-vpc are fine to keep when they are not unique identifiers of a real production account.",
    )
    add_callout(
        document,
        "security",
        "Never paste secret access keys, passwords, or .pem file contents into notes. If a key was exposed, deactivate it in IAM and create a new one. That incident belongs in the troubleshooting table, without the secret value.",
    )


def build_exam_overview(document: Document, assets: dict[str, Path]) -> None:
    heading(document, "SAA-C03 exam map", 1)
    add_body(
        document,
        "The AWS Certified Solutions Architect – Associate exam checks whether you can design AWS solutions that are secure, resilient, high-performing, and cost-aware. "
        "It is aimed at people who design cloud solutions. AWS recommends about one year of hands-on experience. This journal turns your course labs into that kind of experience record.",
    )

    heading(document, "Exam snapshot", 2)
    add_styled_table(
        document,
        ["Topic", "Official outline (confirm before booking)"],
        [
            ["Exam code in this guide", "SAA-C03"],
            ["Role focus", "Solutions architect designing AWS workloads"],
            ["Question count", "65 questions total; 50 scored and 15 unscored. Unscored items are not labeled."],
            ["Time", "130 minutes"],
            ["Question styles", "Multiple choice (one correct answer) and multiple response (two or more correct answers)"],
            ["Scoring", "Scaled score from 100 to 1,000. Passing score is 720. Scoring is compensatory: you pass on the overall score, not on every domain."],
            ["Delivery", "Pearson VUE test center or online proctoring"],
            ["Credential life", "Three years from the pass date, after which recertification is required"],
            ["Published exam fee", "150 USD on the AWS Certification site at the time this framework was written. Confirm the live fee. Foreign-exchange rates and taxes can apply."],
        ],
        col_widths=[2.1, 4.4],
    )
    add_callout(
        document,
        "verify",
        "Exam fee, languages, delivery vendor, and exam code can change. Check the AWS Certified Solutions Architect – Associate page and the current exam guide before you pay or schedule. Unofficial blogs disagree about successor exam codes; this journal follows the official SAA-C03 guide until you confirm otherwise.",
    )

    heading(document, "Content domains and weights", 2)
    add_figure(
        document,
        assets["domains"],
        "Figure 3: SAA-C03 scored-content weights by domain, from the official exam guide.",
        "Horizontal bar chart of SAA-C03 domain weights: Design Secure Architectures 30 percent, Design Resilient Architectures 26 percent, Design High-Performing Architectures 24 percent, and Design Cost-Optimized Architectures 20 percent.",
    )
    add_body(
        document,
        "Each experiment chapter will list one primary domain and any secondary domains. Many labs touch more than one domain. A VPC lab, for example, is a networking exercise and also a security and cost exercise because public subnets, NAT gateways, and security groups appear together.",
    )

    heading(document, "Domain 1: Design Secure Architectures (30%)", 2)
    add_body(
        document,
        "This is the largest scored domain. It asks how identities get access, how applications are isolated on the network, and how data is protected. "
        "In plain language: who can do what, which traffic is allowed, and how information is encrypted and retained.",
    )
    add_styled_table(
        document,
        ["Task", "What you should be able to decide", "Typical lab themes"],
        [
            [
                "1.1 Secure access to AWS resources",
                "Users, groups, roles, policies, federation, multi-account controls, and least privilege, including MFA on privileged access.",
                "IAM users and roles, permission boundaries, Organizations SCPs, IAM Identity Center",
            ],
            [
                "1.2 Secure workloads and applications",
                "VPC design, public versus private subnets, security groups, network ACLs, WAF, Shield, Secrets Manager, and secure connectivity.",
                "VPC labs, security groups, NAT, VPN, endpoints",
            ],
            [
                "1.3 Data security controls",
                "Encryption at rest and in transit, KMS key policies, backups, classification, and lifecycle controls.",
                "S3 encryption, KMS, RDS encryption, ACM certificates",
            ],
        ],
        col_widths=[1.8, 2.5, 2.2],
    )

    heading(document, "Domain 2: Design Resilient Architectures (26%)", 2)
    add_body(
        document,
        "Resilience is the ability to keep serving users when a component, Availability Zone, or Region has a problem. "
        "The exam looks for loosely coupled designs, horizontal scaling, and disaster-recovery choices that match RTO and RPO.",
    )
    add_styled_table(
        document,
        ["Task", "What you should be able to decide", "Typical lab themes"],
        [
            [
                "2.1 Scalable and loosely coupled architectures",
                "Queues, pub/sub, API layers, containers, serverless, multi-tier apps, and independent scaling of components.",
                "SQS, SNS, EventBridge, ALB, Auto Scaling, Lambda, ECS",
            ],
            [
                "2.2 Highly available and fault-tolerant architectures",
                "Multi-AZ placement, failover, backups, Route 53 health checks, and DR patterns such as backup and restore, pilot light, warm standby, and active-active.",
                "RDS Multi-AZ, S3 replication, Route 53 failover, AMI and snapshot restore",
            ],
        ],
        col_widths=[1.8, 2.5, 2.2],
    )

    heading(document, "Domain 3: Design High-Performing Architectures (24%)", 2)
    add_body(
        document,
        "Performance questions ask which storage, compute, database, network, or ingestion option meets a latency, throughput, or scaling need. "
        "The trap is choosing a familiar service when a purpose-built service fits better.",
    )
    add_styled_table(
        document,
        ["Task", "What you should be able to decide", "Typical lab themes"],
        [
            ["3.1 Storage performance and scale", "Object, file, and block storage characteristics and when to use S3, EFS, FSx, or EBS.", "EBS volume types, S3, EFS"],
            ["3.2 Elastic compute", "Instance families, Auto Scaling, Lambda memory sizing, containers, and decoupling so tiers scale separately.", "EC2, Auto Scaling, Lambda, Fargate"],
            ["3.3 Database performance", "Relational versus non-relational, read replicas, caching, and capacity models.", "RDS, Aurora, DynamoDB, ElastiCache"],
            ["3.4 Network performance", "Subnet layout, load balancer choice, CloudFront, Global Accelerator, Direct Connect, and PrivateLink.", "ALB versus NLB, CloudFront, VPC endpoints"],
            ["3.5 Data ingestion and transformation", "Streaming versus batch, Glue, Kinesis family, DataSync, and analytics access patterns.", "Kinesis, Glue, Athena"],
        ],
        col_widths=[1.8, 2.5, 2.2],
    )

    heading(document, "Domain 4: Design Cost-Optimized Architectures (20%)", 2)
    add_body(
        document,
        "Cost questions reward the lowest-price design that still meets the requirement. Right-sizing, storage class, purchase option, and data-transfer path matter as much as the service name.",
    )
    add_styled_table(
        document,
        ["Task", "What you should be able to decide", "Typical lab themes"],
        [
            ["4.1 Cost-optimized storage", "S3 classes, lifecycle rules, EBS volume types, backup retention, and hybrid transfer tools.", "S3 lifecycle, gp3 versus io2, Glacier"],
            ["4.2 Cost-optimized compute", "Spot, Savings Plans, Reserved Instances, Graviton where relevant, Lambda versus EC2, and scaling to zero.", "EC2 purchasing, Auto Scaling, Fargate, Lambda"],
            ["4.3 Cost-optimized databases", "Aurora versus RDS versus DynamoDB, on-demand versus provisioned, and replica strategy.", "RDS sizing, DynamoDB capacity modes"],
            ["4.4 Cost-optimized networks", "NAT gateway placement, VPC endpoints versus NAT, Direct Connect versus VPN, and CloudFront for egress.", "NAT, gateway endpoints, CloudFront"],
        ],
        col_widths=[1.8, 2.5, 2.2],
    )

    heading(document, "Well-Architected connection", 2)
    add_body(
        document,
        "The exam is built around the AWS Well-Architected way of thinking. You do not need to recite pillar names to pass, but you should recognize the questions they represent: "
        "Is access least-privilege? Can the design survive an AZ failure? Will it scale without a rewrite? Is the storage class too expensive for the access pattern? Can you operate and observe it?",
    )
    add_styled_table(
        document,
        ["Pillar", "Question this journal will keep asking after each lab"],
        [
            ["Operational excellence", "How would you observe, change, and recover this design in a real environment?"],
            ["Security", "What identities, network paths, and encryption settings actually protect the workload?"],
            ["Reliability", "What happens if one AZ, one instance, or one database node fails?"],
            ["Performance efficiency", "Is this the right service and size for the access pattern?"],
            ["Cost optimization", "Which resources in this lab start charging as soon as they exist, and how do you shut them off?"],
            ["Sustainability", "Does the design waste idle compute or keep unused storage that does not need to stay hot?"],
        ],
        col_widths=[2.0, 4.5],
    )
    add_callout(
        document,
        "exam",
        "A common exam trap is to pick the most highly available or most expensive option when the scenario gives a recovery-time or budget constraint. Match the design to the stated RTO, RPO, performance, and cost requirements.",
    )


def build_indexes(document: Document) -> None:
    heading(document, "Experiment index", 1)
    add_body(
        document,
        "This index lists every experiment in document order. It is empty of completed labs in version 0.1 because Experiment 1 has not been supplied yet.",
    )
    rows = []
    for exp in EXPERIMENTS:
        rows.append(
            [
                f"EXP{exp.number:02d}",
                exp.name,
                exp.primary_service,
                exp.exam_domains,
                exp.region,
                exp.status,
            ]
        )
    add_styled_table(
        document,
        ["ID", "Experiment name", "Primary service", "Exam domain(s)", "Region", "Status"],
        rows,
        col_widths=[0.8, 1.5, 1.2, 1.2, 1.0, 1.3],
    )

    heading(document, "Index by AWS service", 2)
    add_body(
        document,
        "When experiments are added, this table will group them by the main AWS service so you can revise VPC, IAM, EC2, S3, and other topics as sets. Nothing is listed yet.",
    )
    add_styled_table(
        document,
        ["AWS service", "Experiment IDs", "Notes"],
        [["—", "None yet", "Will be populated after Experiment 1 is approved."]],
        col_widths=[2.0, 1.8, 2.7],
    )

    heading(document, "Index by SAA-C03 domain", 2)
    add_styled_table(
        document,
        ["Domain", "Weight", "Completed experiments", "Coverage note"],
        [
            ["1 Design Secure Architectures", "30%", "None yet", "Highest exam weight. Prioritize IAM, VPC isolation, and encryption labs."],
            ["2 Design Resilient Architectures", "26%", "None yet", "Look for Multi-AZ, queues, Auto Scaling, and DR labs."],
            ["3 Design High-Performing Architectures", "24%", "None yet", "Look for storage, compute sizing, database, and edge labs."],
            ["4 Design Cost-Optimized Architectures", "20%", "None yet", "Every lab should still record what incurs cost and what to delete."],
        ],
        col_widths=[2.3, 0.9, 1.6, 1.7],
    )

    heading(document, "Figure list", 2)
    add_body(
        document,
        "Figure numbers are sequential across the whole document. Experiment screenshots will continue from the next number after these framework figures.",
    )
    add_styled_table(
        document,
        ["Figure", "Caption", "Chapter"],
        [
            ["Figure 1", "Process used to add each experiment to this document.", "How to use this study guide"],
            ["Figure 2", "Architecture diagram conventions used in later experiment chapters.", "Document conventions"],
            ["Figure 3", "SAA-C03 scored-content weights by domain, from the official exam guide.", "SAA-C03 exam map"],
        ],
        col_widths=[1.1, 3.6, 1.8],
    )


def build_experiment_intake(document: Document) -> None:
    heading(document, "Experiment 1: materials required", 1)
    add_callout(
        document,
        "next",
        "Do not skip this page. Send the package below in your next message. Experiment 1 will be written only from that material. If a field is unknown, write unknown rather than guessing.",
        title="Required from you now",
    )
    add_body(
        document,
        "No course notes or screenshots were present in the workspace when this foundation edition was created. "
        "The chapter for Experiment 1 is therefore not written yet. Use this section as the intake form. "
        "A copy-and-fill version also lives at docs/saa-c03/notes/EXP01-intake-template.txt. Redacted screenshots should be saved in docs/saa-c03/screenshots.",
    )

    heading(document, "1. Required fields", 2)
    add_styled_table(
        document,
        ["Field", "What to send", "Example of the kind of answer needed", "Your response"],
        [
            ["Experiment name", "The lab title you used in the course or a short name you prefer", "Create a custom VPC with public and private subnets", ""],
            ["Experiment number in the course", "If the course uses a lab number, include it", "Lab 3, Module 2", ""],
            ["Course section", "Module, week, or video name", "Networking essentials / VPC basics", ""],
            ["AWS Region", "The Region shown in the console during the lab", "ap-south-1 (Mumbai) or us-east-1", ""],
            ["AWS account type", "Personal Free Tier, sandbox, or other. Do not send the account ID", "Personal account, Free Tier", ""],
            ["What you built", "Five to fifteen sentences in your own words", "I created a VPC, two subnets, an internet gateway…", ""],
            ["Why you built it", "The business or exam reason, if the course gave one", "To isolate a web tier from a database tier", ""],
            ["Important settings", "CIDRs, instance types, bucket names, ports, IAM actions, flags", "VPC CIDR 10.0.0.0/16, HTTP 80 from 0.0.0.0/0", ""],
            ["Services used", "Every AWS service you touched", "VPC, EC2, IAM, S3", ""],
            ["Errors encountered", "Exact error text if you still have it, plus how you fixed it", "UnauthorizedOperation on ec2:CreateVpc", ""],
            ["What you want emphasized", "Any topic you found confusing", "Difference between security group and NACL", ""],
        ],
        col_widths=[1.5, 1.7, 1.8, 1.5],
    )

    heading(document, "2. Notes format", 2)
    add_body(document, "Paste notes in the order you performed the work. Rough notes are acceptable. Please include:")
    add_bullet(document, "Each console click path or CLI command, in the order you used it.")
    add_bullet(document, "Names you assigned to VPCs, subnets, instances, buckets, roles, and security groups.")
    add_bullet(document, "CIDR blocks, ports, instance types, storage sizes, and other numeric settings.")
    add_bullet(document, "Whether a resource was public or private.")
    add_bullet(document, "Any default value you left unchanged, if you remember it.")
    add_bullet(document, "How you tested that the lab worked.")
    add_bullet(document, "Whether you already deleted the resources.")

    heading(document, "3. Screenshot package", 2)
    add_body(
        document,
        "Attach every screenshot from the lab. If you already know the step order, rename files before sending using this pattern:",
    )
    add_mixed_paragraph(
        document,
        [("EXP01-Step01-short-description.png", {"bold": True, "size": 11, "color": NAVY})],
    )
    add_body(
        document,
        "If you do not want to rename files, send them in step order and say what each image shows. A simple list such as “Image 1 is the VPC create page” is enough.",
    )
    add_styled_table(
        document,
        ["Good screenshot", "Problematic screenshot"],
        [
            ["The setting being changed is visible and readable.", "Only a desktop wallpaper or an unrelated page is visible."],
            ["Account ID, email, keys, and IPs are hidden.", "The console header still shows a 12-digit account ID or an email."],
            ["Success state is included, not only the blank form.", "The image is cropped so the service name or Region is missing."],
            ["Error dialogs are included when troubleshooting matters.", "The file is so compressed that CIDR text cannot be read."],
        ],
        col_widths=[3.25, 3.25],
    )
    add_callout(
        document,
        "privacy",
        "Before you upload images, hide the 12-digit account ID in the console header, any email address, access keys, secret keys, and public IP addresses. If you cannot blur them, say so and the image will be marked as needing a clearer redacted copy rather than being published with live identifiers.",
    )

    heading(document, "4. Errors and unexpected results", 2)
    add_body(document, "If something failed, include it. Failed steps are valuable exam material. For each issue, send:")
    add_bullet(document, "What you were trying to do.")
    add_bullet(document, "The error message or unexpected result, copied if possible.")
    add_bullet(document, "The service and API action, if shown.")
    add_bullet(document, "What you changed to recover, or “not resolved yet”.")

    heading(document, "5. Information that must not be sent", 2)
    add_bullet(document, "Secret access keys, session tokens, and passwords.")
    add_bullet(document, "Contents of .pem or .ppk private key files.")
    add_bullet(document, "MFA seed codes.")
    add_bullet(document, "Unredacted production account IDs if you are using an employer account.")
    add_bullet(document, "Customer data.")

    heading(document, "6. Missing-item log for Experiment 1", 2)
    add_body(document, "This log will shrink as you send material. It is complete for version 0.1 because nothing has been received yet.")
    add_styled_table(
        document,
        ["Item", "Status", "Effect on the chapter"],
        [
            ["Experiment name", "Missing", "Chapter title cannot be written."],
            ["Course section", "Missing", "Exam mapping can only be inferred after services are known."],
            ["Learner notes", "Missing", "Procedure, resource table, and architecture cannot be written."],
            ["AWS Region", "Missing", "Prerequisites and service-availability notes cannot be completed."],
            ["Screenshots", "Missing", "No step images can be inserted."],
            ["Errors encountered", "Missing", "Troubleshooting table will stay generic until you send errors, or will say none reported."],
            ["Resource names and CIDRs", "Missing", "Configuration table will use [Verification required] if notes are incomplete."],
            ["Cleanup status", "Missing", "Cleanup chapter will assume resources may still exist until you say otherwise."],
        ],
        col_widths=[2.2, 1.3, 3.0],
    )

    heading(document, "7. Confirmation you will be asked to make later", 2)
    add_body(
        document,
        "After a draft chapter exists, you will be asked to confirm values such as Region, CIDR blocks, instance types, ports, bucket names, and whether encryption or Multi-AZ was enabled. "
        "Until you confirm them, those values remain marked [Verification required].",
    )

    heading(document, "Experiment 1 chapter status", 2)
    add_styled_table(
        document,
        ["Section", "Status"],
        [
            ["1. Experiment Overview", "Not started — waiting for notes"],
            ["2. SAA-C03 Exam Mapping", "Not started — waiting for notes"],
            ["3. Learning Objectives", "Not started — waiting for notes"],
            ["4. Architecture", "Not started — waiting for notes"],
            ["5. Prerequisites", "Not started — waiting for notes"],
            ["6. Resources and Configuration", "Not started — waiting for notes"],
            ["7. Step-by-Step Procedure", "Not started — waiting for screenshots and notes"],
            ["8. Verification", "Not started — waiting for notes"],
            ["9. Troubleshooting", "Not started — waiting for errors or a statement that none occurred"],
            ["10. Security Best Practices", "Not started — waiting for notes"],
            ["11. Cost Considerations", "Not started — waiting for notes"],
            ["12. Cleanup Procedure", "Not started — waiting for notes"],
            ["13. Important SAA-C03 Exam Points", "Not started — waiting for notes"],
            ["14. Key Terms", "Not started — waiting for notes"],
            ["15. Knowledge Check", "Not started — waiting for notes"],
            ["16. Experiment Summary", "Not started — waiting for notes"],
        ],
        col_widths=[3.2, 3.3],
    )


def build_progress(document: Document) -> None:
    heading(document, "Progress tracker and revision checkpoints", 1)
    add_body(
        document,
        "After every five completed experiments, this part will receive a checkpoint: what was covered, which exam domains are still thin, which screenshots still need clearer copies, and which cleanup items remain open.",
    )

    heading(document, "Overall progress", 2)
    add_styled_table(
        document,
        ["Measure", "Current value", "Target for the next checkpoint"],
        [
            ["Document version", "0.1 Foundation", "0.2 after Experiment 1 is approved"],
            ["Experiments completed", "0", "5 completed experiments trigger Checkpoint A"],
            ["Experiments awaiting materials", "1 (Experiment 1)", "0 awaiting packages"],
            ["Figures in document", "3 framework figures", "Framework figures plus Experiment 1 screenshots"],
            ["Glossary terms", "Starter set in this edition", "Grow with each experiment’s Key Terms table"],
            ["Known [Verification required] items", "Exam live-code confirmation; all Experiment 1 fields", "Only items you still cannot confirm"],
        ],
        col_widths=[2.3, 2.1, 2.1],
    )

    heading(document, "Checkpoint A — after Experiments 1 to 5", 2)
    add_body(document, "This checkpoint will be written when five experiment chapters are approved. Planned contents:")
    add_bullet(document, "List of the five experiments, services, and domains.")
    add_bullet(document, "Domain coverage compared with exam weights.")
    add_bullet(document, "Repeated mistakes or exam traps.")
    add_bullet(document, "Open cleanup or cost risks.")
    add_bullet(document, "Screenshot quality issues still open.")
    add_bullet(document, "Suggested focus for Experiments 6 to 10.")
    add_callout(
        document,
        "note",
        "Checkpoints B and C will be added after Experiments 10 and 15. They are not created in advance as empty chapters.",
    )


def build_glossary(document: Document) -> None:
    heading(document, "Glossary", 1)
    add_body(
        document,
        "These starter definitions use simple language. Experiment chapters will add terms that appear in that lab. If a definition later disagrees with an official AWS glossary term, the official meaning wins and this table will be corrected.",
    )
    add_styled_table(
        document,
        ["Term", "Simple explanation"],
        [
            ["Availability Zone (AZ)", "An isolated location inside a Region, with its own buildings and power. Designing across two or more AZs is the usual way to survive one data-center failure."],
            ["Region", "A geographic area that contains multiple Availability Zones. You choose a Region for latency, data-residency, and service availability."],
            ["VPC", "Your private network space in AWS. You control IP ranges, subnets, route tables, and gateways."],
            ["Subnet", "A slice of VPC IP space in one Availability Zone. Public subnets have a route to an internet gateway. Private subnets do not."],
            ["Internet gateway", "The VPC component that lets resources with public IP addresses send and receive traffic from the internet."],
            ["NAT gateway", "A managed translator that lets private-subnet resources start outbound internet connections without being reachable from the internet."],
            ["Route table", "The set of rules that decides where subnet traffic goes next, such as local, internet gateway, NAT gateway, or peering."],
            ["Security group", "A stateful virtual firewall attached to an elastic network interface. Return traffic is allowed automatically."],
            ["Network ACL", "A stateless subnet firewall. Inbound and outbound rules are evaluated separately and are easy to misconfigure."],
            ["IAM user", "A long-lived identity in an AWS account, usually for a person or a static program. Prefer roles over long-lived access keys."],
            ["IAM role", "An identity that a user or service can assume. It has permissions but no long-term password of its own."],
            ["IAM policy", "A JSON document that allows or denies actions on resources. Keep the actions and resources as small as the job allows."],
            ["Least privilege", "Give only the permissions required for the task, not AdministratorAccess “to save time” in a real design."],
            ["Shared responsibility model", "AWS secures the cloud infrastructure. You secure what you put in the cloud, including identities, data, and network configuration. The split shifts for managed services."],
            ["Encryption at rest", "Data is stored in unreadable form, usually with KMS-managed keys, until an authorized service decrypts it."],
            ["Encryption in transit", "Data is protected while it moves, commonly with TLS certificates from ACM."],
            ["S3", "Object storage for files of any common size. Durability is extremely high. Storage classes trade retrieval time against cost."],
            ["EBS", "Block storage attached to an EC2 instance in the same AZ. Choose volume type based on IOPS, throughput, and cost."],
            ["EFS", "A managed NFS file system that multiple instances can mount, including across AZs in a Region."],
            ["EC2", "Virtual servers you manage at the operating-system level. You choose the instance family, size, purchasing option, and networking."],
            ["Auto Scaling group", "A fleet that adds or removes EC2 instances from a launch template based on demand or schedule."],
            ["Application Load Balancer", "A Layer 7 load balancer. It can route on host, path, and headers and is the usual front door for HTTP applications."],
            ["Network Load Balancer", "A Layer 4 load balancer for extreme performance, static IP needs, or non-HTTP protocols."],
            ["RDS", "Managed relational databases. Multi-AZ is a standby for failover. A read replica is for read scaling and is not the same thing."],
            ["Aurora", "A MySQL- and PostgreSQL-compatible AWS database engine with a shared storage design and typically faster failover than classic RDS."],
            ["DynamoDB", "A managed key-value and document database. You design around the partition key. It is not a drop-in replacement for SQL joins."],
            ["Lambda", "Run code without managing servers. You pay for invocation and duration. It has time, memory, and payload limits."],
            ["CloudFront", "A content delivery network that caches content at edge locations closer to users and can also protect origins."],
            ["Route 53", "AWS DNS. Routing policies include simple, failover, latency, geolocation, geoproximity, weighted, and multi-value."],
            ["KMS", "Creates and controls encryption keys. Key policies and IAM policies both matter. Know customer managed keys versus AWS owned keys at a high level."],
            ["CloudTrail", "Records API activity for audit. This is who did what, not the same as CloudWatch metrics."],
            ["CloudWatch", "Metrics, logs, and alarms for operations. Use it to know how a workload is behaving."],
            ["RTO", "Recovery time objective: how quickly the service must return after a disaster."],
            ["RPO", "Recovery point objective: how much data you can afford to lose, measured in time since the last good copy."],
            ["Multi-AZ", "A design that keeps a standby or copies in another Availability Zone so a single AZ outage is survivable."],
            ["Stateless versus stateful", "A stateless tier can be replaced at any time. Session or data that must survive belongs in a shared store, not on one instance disk."],
        ],
        col_widths=[2.0, 4.5],
    )


def build_appendices(document: Document) -> None:
    heading(document, "Appendix A: Experiment chapter template", 1)
    add_body(
        document,
        "The following template is the empty shell copied for every new experiment. Placeholder sentences show what belongs in each section. They are not an invented lab.",
    )

    heading(document, "Experiment [Number]: [Experiment Name]", 2)
    heading(document, "1. Experiment Overview", 3)
    add_body(document, "State what is being built and why that pattern appears in real AWS environments. Use only facts from the notes.")
    heading(document, "2. SAA-C03 Exam Mapping", 3)
    add_body(document, "List exam domain, AWS services, important exam concepts, and difficulty: Beginner, Intermediate, or Advanced.")
    heading(document, "3. Learning Objectives", 3)
    add_body(document, "List skills the learner can perform after the lab, written as observable actions.")
    heading(document, "4. Architecture", 3)
    add_body(document, "Insert a simple diagram of users, services, data flow, network boundaries, and security controls, then explain it in a short paragraph.")
    heading(document, "5. Prerequisites", 3)
    add_body(document, "Record required permissions, services, Region, files or tools, estimated time, and possible charges. Do not invent a price.")
    heading(document, "6. Resources and Configuration", 3)
    add_body(document, "Fill the table: Resource | Name | Important Configuration | Purpose.")
    heading(document, "7. Step-by-Step Procedure", 3)
    add_body(
        document,
        "For each step use Action, Why this step matters, Expected result, and Screenshot. Place the matching image immediately after the instruction, then describe what the screenshot shows.",
    )
    heading(document, "8. Verification", 3)
    add_body(document, "Include console checks, useful CLI commands, expected output, connectivity or application tests, and a permission or security check.")
    heading(document, "9. Troubleshooting", 3)
    add_body(document, "Fill Problem/Error | Likely Cause | Recommended Solution, including errors the learner actually hit.")
    heading(document, "10. Security Best Practices", 3)
    add_body(document, "Cover least privilege, encryption, security groups and NACLs, public versus private access, secrets, logging, and Multi-AZ or DR where relevant.")
    heading(document, "11. Cost Considerations", 3)
    add_body(document, "Name chargeable resources, Free Tier eligibility only when applicable and not as a guarantee, how to avoid idle cost, and what must be deleted. State that AWS pricing varies by Region and over time.")
    heading(document, "12. Cleanup Procedure", 3)
    add_body(document, "Delete only resources created in the experiment, in dependency order, then give a final check that nothing billable remains.")
    heading(document, "13. Important SAA-C03 Exam Points", 3)
    add_body(document, "Highlight likely exam facts, commonly confused services, trade-offs, traps, and when not to use a service.")
    heading(document, "14. Key Terms", 3)
    add_body(document, "Add Term | Simple Explanation for words introduced in the lab.")
    heading(document, "15. Knowledge Check", 3)
    add_body(document, "Write five original SAA-C03-style multiple-choice questions with four options. Place answers and short explanations after the question set. Do not copy vendor exam questions.")
    heading(document, "16. Experiment Summary", 3)
    add_body(document, "Finish with five to eight concise points.")

    heading(document, "Appendix B: Screenshot naming and quality checklist", 1)
    add_styled_table(
        document,
        ["Check", "Pass condition"],
        [
            ["File name", "EXP01-Step01-VPC-Creation.png style, ASCII, no spaces."],
            ["Step match", "Image sits under the step that created that console view."],
            ["Readability", "Setting names and values can be read at the inserted size."],
            ["Redaction", "Account IDs, emails, keys, and IPs are hidden."],
            ["No value edits", "CIDRs, instance types, and other settings were not painted over except for redaction boxes."],
            ["Alt text", "The image has a short description of the technical content."],
            ["Caption", "Figure N: action-focused caption under the image."],
            ["Unclear image", "Replaced, or marked [A clearer screenshot is required]."],
        ],
        col_widths=[1.8, 4.7],
    )

    heading(document, "Appendix C: Generic cost and cleanup hygiene", 1)
    add_body(
        document,
        "This appendix is not a substitute for an experiment’s cleanup section. It is a safety net for any lab. Always prefer the chapter-specific order when it exists, because some resources cannot be deleted until dependents are gone.",
    )
    add_callout(
        document,
        "cost",
        "AWS pricing depends on Region, volume, and the date you look it up. This journal will not invent a currency amount. If a resource can generate a charge even while idle, that fact will be stated in the experiment’s cost section after the resource list is known.",
    )
    add_body(document, "Resources that commonly continue charging if left running include:")
    add_bullet(document, "NAT gateways, Elastic Load Balancers, and VPC endpoints of the interface type.")
    add_bullet(document, "EC2 instances, EBS volumes, and unattached Elastic IP addresses.")
    add_bullet(document, "RDS, Aurora, ElastiCache, and OpenSearch domains.")
    add_bullet(document, "NAT and data-transfer paths that stay idle but provisioned.")
    add_bullet(document, "CloudWatch logs stored for a long retention period, and S3 objects you forget to delete.")
    add_body(document, "Typical deletion order, from dependents to foundations:")
    add_numbered(document, "Stop or delete compute (instances, functions, containers) and load balancers.")
    add_numbered(document, "Delete databases, file systems, and other data stores, including snapshots you created for the lab if they are not needed.")
    add_numbered(document, "Release Elastic IPs and delete NAT gateways.")
    add_numbered(document, "Delete custom security groups, route table associations, subnets, internet gateways, and the VPC.")
    add_numbered(document, "Delete lab IAM users, access keys, and roles that were created only for the experiment.")
    add_numbered(document, "Empty and delete lab S3 buckets.")
    add_numbered(document, "Check Billing and Cost Management and the Resource Groups / tag filter for the lab prefix.")

    heading(document, "Appendix D: Refresh the table of contents in Microsoft Word", 1)
    add_body(
        document,
        "The contents page is a Word field. It updates from Heading 1 to Heading 3. On first open, allow Word to update fields. If the list still shows the placeholder sentence, use one of these methods.",
    )
    add_numbered(document, "Click inside the contents list.")
    add_numbered(document, "Press F9, or right-click and choose Update Field.")
    add_numbered(document, "Choose Update entire table.")
    add_numbered(document, "On the References tab, choose Update Table.")
    add_body(
        document,
        "If you export to PDF from Word, update the contents field immediately before export so page numbers in the PDF match the printed pages.",
    )

    heading(document, "Appendix E: Knowledge-check writing rules", 1)
    add_body(document, "Questions added after each experiment will follow these rules so they stay useful for SAA-C03 practice:")
    add_bullet(document, "Five questions per experiment, each with four options.")
    add_bullet(document, "One best answer unless the question is clearly a multi-response style. Version 0.1 will use single-best-answer items unless a lab naturally needs two correct actions.")
    add_bullet(document, "Stem based on an architecture choice, not trivia such as a console button color.")
    add_bullet(document, "Distractors will be plausible AWS services, which is how the real exam is written.")
    add_bullet(document, "Answers appear after the full question set so you can test yourself.")
    add_bullet(document, "No copied exam-dump items and no verbatim course quiz items.")

    heading(document, "Closing status for version 0.1", 1)
    add_body(
        document,
        "The professional document structure is in place: title page, automatic contents field, exam domain map, indexes, figure conventions, privacy rules, glossary, and the 16-section experiment template. "
        "The next document update should be Experiment 1, written from your notes and screenshots.",
    )
    add_callout(
        document,
        "next",
        "Reply with Experiment 1’s name, course section, notes, AWS Region, screenshots, and any errors. If screenshots are not ready, send the notes first and mark missing images. The chapter will still list every missing screenshot instead of filling the gap with invented steps.",
    )


def set_core_properties(document: Document) -> None:
    props = document.core_properties
    props.title = DOCUMENT["title"]
    props.subject = "Personal SAA-C03 practical-lab journal and exam study companion"
    props.author = DOCUMENT["learner"]
    props.category = "Study guide"
    props.comments = (
        "Personal study notes. Not official AWS training. Version "
        + DOCUMENT["version"]
        + ". Experiment chapters are added only after notes or screenshots are supplied."
    )
    props.keywords = "AWS, SAA-C03, Solutions Architect Associate, practical labs, study guide"


def build() -> Path:
    assets = generate_all()
    document = Document()
    set_document_defaults(document)
    set_core_properties(document)
    add_header_and_footer(document.sections[0], DOCUMENT["short_title"], "Version " + DOCUMENT["version"])
    set_update_fields_on_open(document)

    build_cover(document, assets)
    add_page_break(document)

    heading(document, "Contents", 1)
    add_toc(document)
    add_page_break(document)

    build_inside_cover(document)
    build_how_to_use(document, assets)
    build_conventions(document, assets)
    build_privacy(document)
    build_exam_overview(document, assets)
    build_indexes(document)
    build_experiment_intake(document)
    build_progress(document)
    build_glossary(document)
    build_appendices(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    set_update_fields_on_open(document)
    document.save(OUTPUT)
    patch_docx_update_fields(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes)")
