"""
Shared Word theme, styles, and helpers for the SAA-C03 study guide.

Keep visual rules here so every future experiment chapter uses the same
headings, tables, callouts, captions, headers, and footers.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


# AWS-inspired professional palette (not an official AWS template).
NAVY = RGBColor(0x23, 0x2F, 0x3E)
NAVY_HEX = "232F3E"
ORANGE = RGBColor(0xFF, 0x99, 0x00)
ORANGE_HEX = "FF9900"
TEAL = RGBColor(0x1B, 0x6B, 0x93)
TEAL_HEX = "1B6B93"
SLATE = RGBColor(0x54, 0x5B, 0x64)
BODY = RGBColor(0x2B, 0x2B, 0x2B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WHITE_HEX = "FFFFFF"
LIGHT_NAVY_HEX = "E8EEF2"
ROW_ALT_HEX = "F4F7F9"
AMBER_HEX = "FFF4D6"
AMBER_BORDER_HEX = "C48500"
GREEN_HEX = "E7F6EC"
GREEN_BORDER_HEX = "1E8E3E"
RED_HEX = "FDECEC"
RED_BORDER_HEX = "D13212"
BLUE_HEX = "E8F4FC"
BLUE_BORDER_HEX = "1B6B93"
ORANGE_TINT_HEX = "FFF6E8"
GRAY_HEX = "F2F3F3"
GRAY_BORDER_HEX = "687078"

FONT_NAME = "Calibri"
FONT_HEADING = "Calibri"

CALLOUT_STYLES = {
    "exam": {
        "fill": ORANGE_TINT_HEX,
        "accent": ORANGE_HEX,
        "label": "Exam tip",
        "label_color": RGBColor(0x8A, 0x5A, 0x00),
    },
    "security": {
        "fill": RED_HEX,
        "accent": RED_BORDER_HEX,
        "label": "Security",
        "label_color": RGBColor(0xA0, 0x1D, 0x12),
    },
    "cost": {
        "fill": GREEN_HEX,
        "accent": GREEN_BORDER_HEX,
        "label": "Cost",
        "label_color": RGBColor(0x14, 0x6C, 0x2E),
    },
    "verify": {
        "fill": AMBER_HEX,
        "accent": AMBER_BORDER_HEX,
        "label": "Verification required",
        "label_color": RGBColor(0x8A, 0x5A, 0x00),
    },
    "privacy": {
        "fill": GRAY_HEX,
        "accent": GRAY_BORDER_HEX,
        "label": "Privacy",
        "label_color": SLATE,
    },
    "next": {
        "fill": BLUE_HEX,
        "accent": BLUE_BORDER_HEX,
        "label": "Next action",
        "label_color": TEAL,
    },
    "note": {
        "fill": LIGHT_NAVY_HEX,
        "accent": NAVY_HEX,
        "label": "Note",
        "label_color": NAVY,
    },
    "warning": {
        "fill": AMBER_HEX,
        "accent": RED_BORDER_HEX,
        "label": "Warning",
        "label_color": RGBColor(0xA0, 0x1D, 0x12),
    },
}


def _element(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(key), value)
    return el


def set_run_font(run, size_pt: float, *, bold=False, italic=False, color=None, name=FONT_NAME):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell: _Cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = _element("w:shd", **{"w:fill": hex_color, "w:val": "clear", "w:color": "auto"})
    tcPr.append(shd)


def set_cell_margins(cell: _Cell, top=80, bottom=80, left=100, right=100) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = _element("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = _element(f"w:{edge}", **{"w:w": str(value), "w:type": "dxa"})
        tcMar.append(node)
    for child in list(tcPr):
        if child.tag == qn("w:tcMar"):
            tcPr.remove(child)
    tcPr.append(tcMar)


def set_cell_border(cell: _Cell, **sides) -> None:
    """
    sides example: top={"sz": "4", "color": "232F3E", "val": "single"}
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = _element("w:tcBorders")
        tcPr.append(tcBorders)
    for edge, spec in sides.items():
        element = _element(f"w:{edge}")
        for attr, value in spec.items():
            element.set(qn(f"w:{attr}"), str(value))
        existing = tcBorders.find(qn(f"w:{edge}"))
        if existing is not None:
            tcBorders.remove(existing)
        tcBorders.append(element)


def set_table_borders(table: Table, color: str = "C5CDD3", sz: str = "4") -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = _element("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = _element(f"w:{edge}", **{"w:val": "single", "w:sz": sz, "w:space": "0", "w:color": color})
        old = borders.find(qn(f"w:{edge}"))
        if old is not None:
            borders.remove(old)
        borders.append(el)


def prevent_row_split(row) -> None:
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = _element("w:cantSplit")
    trPr.append(cant)


def set_keep_with_next(paragraph: Paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(_element("w:keepNext"))


def set_keep_together(paragraph: Paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.append(_element("w:keepLines"))


def disable_autofit(table: Table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = _element("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_table_width(table: Table, width_dxa: int) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = _element("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(width_dxa))
    tblW.set(qn("w:type"), "dxa")


def cell_paragraph(cell: _Cell) -> Paragraph:
    if cell.paragraphs:
        return cell.paragraphs[0]
    return cell.add_paragraph()


def write_cell(
    cell: _Cell,
    text: str,
    *,
    bold=False,
    italic=False,
    size=10.5,
    color=BODY,
    align="left",
    fill=None,
    font_name=FONT_NAME,
) -> None:
    if fill:
        shade_cell(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    paragraph = cell_paragraph(cell)
    paragraph.clear()
    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold, italic=italic, color=color, name=font_name)


def add_horizontal_line(paragraph: Paragraph, color: str = NAVY_HEX, size: str = "12") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = _element("w:pBdr")
    bottom = _element(
        "w:bottom",
        **{"w:val": "single", "w:sz": size, "w:space": "1", "w:color": color},
    )
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_field_run(paragraph: Paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    r = run._r
    r.append(_element("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = _element("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    r.append(instr)
    r.append(_element("w:fldChar", **{"w:fldCharType": "separate"}))
    visible = paragraph.add_run("1")
    set_run_font(visible, 9, color=SLATE)
    end_run = paragraph.add_run()
    end_run._r.append(_element("w:fldChar", **{"w:fldCharType": "end"}))


def set_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is not None:
        settings.remove(existing)
    settings.append(_element("w:updateFields", **{"w:val": "true"}))


def patch_docx_update_fields(docx_path: Path) -> None:
    """Guarantee Word refreshes TOC and page-number fields on first open."""
    from zipfile import ZIP_DEFLATED, ZipFile

    tmp_path = docx_path.with_suffix(".updating.docx")
    with ZipFile(docx_path, "r") as zin, ZipFile(tmp_path, "w", compression=ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                text = data.decode("utf-8")
                if "w:updateFields" not in text:
                    text = text.replace(
                        "</w:settings>",
                        '<w:updateFields w:val="true"/></w:settings>',
                    )
                    data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp_path.replace(docx_path)


def set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.05)
    section.bottom_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    section.different_first_page_header_footer = True

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _element("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), FONT_NAME)

    heading_specs = {
        "Heading 1": (18, NAVY, 16, 8, True),
        "Heading 2": (14, NAVY, 14, 6, True),
        "Heading 3": (12, TEAL, 12, 4, True),
        "Heading 4": (11, RGBColor(0x3D, 0x48, 0x56), 10, 4, True),
    }
    for name, (size, color, before, after, bold) in heading_specs.items():
        style = styles[name]
        style.font.name = FONT_HEADING
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.08
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = _element("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), FONT_HEADING)

    caption = styles["Caption"]
    caption.font.name = FONT_NAME
    caption.font.size = Pt(10)
    caption.font.italic = True
    caption.font.color.rgb = SLATE
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(12)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_header_and_footer(section, title: str, version: str) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    run = hp.add_run(title)
    set_run_font(run, 9, bold=True, color=NAVY)
    add_horizontal_line(hp, ORANGE_HEX, "12")

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].clear()

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_horizontal_line(fp, NAVY_HEX, "8")
    left = fp.add_run("Personal study notes  ·  Not official AWS training  ·  ")
    set_run_font(left, 8.5, color=SLATE)
    ver = fp.add_run(version)
    set_run_font(ver, 8.5, color=SLATE)
    spacer = fp.add_run("    ")
    set_run_font(spacer, 8.5, color=SLATE)
    page_label = fp.add_run("Page ")
    set_run_font(page_label, 8.5, color=SLATE)
    add_field_run(fp, " PAGE ")
    of_run = fp.add_run(" of ")
    set_run_font(of_run, 8.5, color=SLATE)
    add_field_run(fp, " NUMPAGES ")

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    ffp = first_footer.paragraphs[0]
    ffp.clear()
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = ffp.add_run("Personal study journal  ·  Hide account IDs, keys, emails, and IP addresses before sharing")
    set_run_font(note, 8.5, color=SLATE)


def add_toc(document: Document) -> None:
    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    run = intro.add_run(
        "Word builds this table of contents from the document headings. "
        "When you first open the file, choose Yes if Word asks to update fields. "
        "You can also refresh it later with References → Update Table."
    )
    set_run_font(run, 10.5, italic=True, color=SLATE)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    r = run._r
    r.append(_element("w:fldChar", **{"w:fldCharType": "begin"}))
    instr = _element("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    r.append(instr)
    r.append(_element("w:fldChar", **{"w:fldCharType": "separate"}))
    placeholder = paragraph.add_run(
        "Right-click this field and choose Update Field to display headings."
    )
    set_run_font(placeholder, 11, italic=True, color=SLATE)
    end = paragraph.add_run()
    end._r.append(_element("w:fldChar", **{"w:fldCharType": "end"}))


def add_body(document: Document, text: str, *, space_after=8) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    set_run_font(run, 11, color=BODY)
    return paragraph


def add_mixed_paragraph(document: Document, parts: list[tuple[str, dict]], *, space_after=8) -> Paragraph:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    for text, opts in parts:
        run = paragraph.add_run(text)
        set_run_font(
            run,
            opts.get("size", 11),
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
            color=opts.get("color", BODY),
        )
    return paragraph


def add_bullet(document: Document, text: str, *, level=0) -> Paragraph:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.clear()
    paragraph.paragraph_format.left_indent = Inches(0.25 + (0.25 * level))
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, 11, color=BODY)
    return paragraph


def add_numbered(document: Document, text: str) -> Paragraph:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, 11, color=BODY)
    return paragraph


def add_callout(document: Document, kind: str, body: str, title: str | None = None) -> Table:
    spec = CALLOUT_STYLES[kind]
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, spec["fill"])
    set_cell_margins(cell, top=80, bottom=90, left=140, right=140)
    set_cell_border(
        cell,
        top={"sz": "4", "val": "single", "color": spec["accent"]},
        bottom={"sz": "4", "val": "single", "color": spec["accent"]},
        right={"sz": "4", "val": "single", "color": spec["accent"]},
        left={"sz": "24", "val": "single", "color": spec["accent"]},
    )
    p = cell_paragraph(cell)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    label = p.add_run((title or spec["label"]).upper())
    set_run_font(label, 9, bold=True, color=spec["label_color"])
    body_p = cell.add_paragraph()
    body_p.paragraph_format.space_after = Pt(0)
    body_p.paragraph_format.space_before = Pt(2)
    body_run = body_p.add_run(body)
    set_run_font(body_run, 10.5, color=BODY)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(8)
    after.paragraph_format.space_before = Pt(0)
    return table


def add_styled_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    *,
    col_widths: list[float] | None = None,
    header_fill: str = NAVY_HEX,
) -> Table:
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    usable = Inches(6.5)
    if col_widths:
        widths = [Inches(w) for w in col_widths]
    else:
        widths = [usable / len(headers)] * len(headers)
    disable_autofit(table)
    set_table_width(table, int(sum(w.twips for w in widths)))
    set_table_borders(table, "C5CDD3", "4")

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        write_cell(cell, header, bold=True, size=10, color=WHITE, fill=header_fill)
        cell.width = widths[idx]
    prevent_row_split(table.rows[0])

    for r_idx, row in enumerate(rows):
        fill = WHITE_HEX if r_idx % 2 == 0 else ROW_ALT_HEX
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            write_cell(cell, value, size=10, color=BODY, fill=fill)
            cell.width = widths[c_idx]
        prevent_row_split(table.rows[r_idx + 1])

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(2)
    spacer.paragraph_format.space_after = Pt(10)
    return table


def set_picture_alt_text(inline_shape, alt_text: str, title: str = "") -> None:
    inline = inline_shape._inline
    docPr = inline.docPr
    docPr.set("descr", alt_text)
    if title:
        docPr.set("name", title)
    cNvPr = inline.graphic.graphicData.pic.nvPicPr.cNvPr
    cNvPr.set("descr", alt_text)


def add_figure(
    document: Document,
    image_path: Path,
    caption: str,
    alt_text: str,
    *,
    max_width_in=6.3,
    max_height_in=4.4,
) -> None:
    from PIL import Image

    with Image.open(image_path) as img:
        width_px, height_px = img.size
    aspect = height_px / float(width_px)
    width = max_width_in
    height = width * aspect
    if height > max_height_in:
        height = max_height_in
        width = height / aspect

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(6)
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    set_picture_alt_text(inline, alt_text, title=caption)

    cap = document.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, 10, italic=True, color=SLATE)


def add_page_break(document: Document) -> None:
    document.add_page_break()
